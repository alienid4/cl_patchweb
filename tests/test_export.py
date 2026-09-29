import datetime as dt
import io

from openpyxl import load_workbook

from webvuln import export, importer
from webvuln.schemas import FindingIn, ImportIn

TODAY = dt.date(2026, 5, 10)


def _load(session):
    findings = [
        FindingIn(
            sheet_key="主機弱點", plugin_id="p1", host="h1", severity="High",
            owner="玄慈", remediation_due="2026-06-01", close_status="未結案",
            raw={"主機": "h1", "弱點名稱": "OpenSSL", "負責人": "玄慈", "備註": ""},
        ),
        FindingIn(
            sheet_key="主機弱點", plugin_id="p2", host="h2", severity="Low",
            owner="", remediation_due=None, close_status="未結案",
            raw={"主機": "h2", "弱點名稱": "Apache"},  # 缺「負責人」「備註」鍵
        ),
    ]
    importer.create_batch(session, ImportIn(
        findings=findings,
        sheet_columns={"主機弱點": ["主機", "弱點名稱", "負責人", "備註"]},
    ))


def test_export_columns_1to1_and_missing(session):
    _load(session)
    wb, fname = export.build_workbook(session, today=TODAY)
    assert fname.endswith(".xlsx")
    assert "管理摘要" in wb.sheetnames
    assert "主機弱點" in wb.sheetnames

    ws = wb["主機弱點"]
    header = [c.value for c in ws[1]]
    assert header == ["主機", "弱點名稱", "負責人", "備註"]      # 1:1 欄序

    r2 = [c.value for c in ws[2]]
    assert r2 == ["h1", "OpenSSL", "玄慈", export.MISSING]      # 空字串→無原始資料
    r3 = [c.value for c in ws[3]]
    assert r3 == ["h2", "Apache", export.MISSING, export.MISSING]  # 缺鍵→無原始資料


def test_export_summary_and_bytes(session):
    _load(session)
    wb, _ = export.build_workbook(session, today=TODAY)
    ws = wb["管理摘要"]
    metrics = {row[0].value: row[1].value for row in ws.iter_rows(min_row=2)}
    assert metrics["未結案"] == 2
    assert metrics["無負責人"] == 1
    assert metrics["無到期日"] == 1

    data = export.to_bytes(wb)
    assert data[:2] == b"PK"  # xlsx = zip
    load_workbook(io.BytesIO(data))  # 可被重新開啟


def test_export_no_data(session):
    wb, fname = export.build_workbook(session, today=TODAY)
    assert wb["管理摘要"]["A2"].value == "狀態"
    assert fname.endswith(".xlsx")
