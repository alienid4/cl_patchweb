"""清單完整匯出 xlsx：原檔欄序、日期與數字型別還原、承辦欄接最後（2026-10-12：CSV 貼回原檔會走樣）。"""
import datetime as dt
import io

from openpyxl import load_workbook


def test_export_rows_xlsx_types_and_order(client):
    raw = {"Plugin ID": "242168", "Host": "10.30.19.5", "Port": "443", "修補期限": "2026/07/19",
           "Description": "line1\nline2", "編號": "0012"}
    fs = [dict(sheet_key="s1", plugin_id="242168", name="v", host="10.30.19.5", severity="High", department="A",
               owner="x", remediation_due="2026-07-19", close_status="未結案", raw=raw)]
    r = client.post("/api/import", json={"source_file": "t.xlsx", "findings": fs,
                                         "sheet_columns": {"s1": ["Plugin ID", "Host", "Port", "修補期限", "Description", "編號"]}})
    assert r.status_code == 200
    fid = client.get("/api/findings").json()[0]["id"]
    client.post(f"/api/findings/{fid}/overlay", json={"set_progress": True, "progress": "要申請展延"})
    x = client.post("/api/export/rows", json={"ids": [fid], "title": "測試清單"})
    assert x.status_code == 200
    ws = load_workbook(io.BytesIO(x.content)).worksheets[0]
    head = [c.value for c in ws[1]]
    assert head[:6] == ["Plugin ID", "Host", "Port", "修補期限", "Description", "編號"]
    assert head[6] == "承辦_處理進度"
    row = [c.value for c in ws[2]]
    assert row[0] == 242168 and row[2] == 443                     # 數字
    assert row[3] == dt.datetime(2026, 7, 19)                     # 真日期
    assert ws.cell(row=2, column=4).number_format == "yyyy/m/d"
    assert row[5] == "0012"                                       # 前導 0 保留文字
    assert row[6] == "要申請展延"
    assert not ws.cell(row=2, column=5).alignment.wrap_text
