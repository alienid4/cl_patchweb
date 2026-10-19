"""原封匯出：把某快照 1:1 還原成 xlsx。

規則（使用者定調）：
- 來源 Excel 有幾欄，匯出就有幾欄，欄名／欄序一模一樣（靠 sheet_columns 記的欄序）。
- 某格在來源沒有值 → 寫「無原始資料」，這樣一看就知道是「來源缺」不是「系統(AIX)出錯」。
- 另加一張「管理摘要」分頁放關鍵數字，方便主管直接看。
匯出走 raw（整列原始欄位）。工具用的正規化欄不參與原封還原，只用於摘要。
"""
from __future__ import annotations

import datetime as dt
import io
import re
from typing import Optional

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from sqlalchemy import select
from sqlalchemy.orm import Session

from . import query
from .models import Finding, ImportBatch, SheetColumns

MISSING = "無原始資料"

# 沒有 sheet_columns（W1 之前的批）時，用 raw 的鍵推欄序；再不行用這組工具欄墊底
_FALLBACK_COLS = ["sheet_key", "plugin_id", "name", "host", "severity",
                  "department", "owner", "remediation_due", "close_status", "remark"]


def _pick_batch(session: Session, batch_id: Optional[int]) -> Optional[ImportBatch]:
    if batch_id is not None:
        return session.get(ImportBatch, batch_id)
    return query.latest_batch(session)


def _cell(raw: Optional[dict], col: str):
    """取原值；缺鍵或空值 → 無原始資料（區分來源缺 vs 系統錯）。"""
    if not raw or col not in raw:
        return MISSING
    v = raw.get(col)
    if v is None or (isinstance(v, str) and v.strip() == ""):
        return MISSING
    return v


def _safe_sheet_name(name: str, used: set) -> str:
    n = re.sub(r"[\[\]\:\*\?\/\\]", "_", (name or "sheet"))[:31] or "sheet"
    base, i = n, 1
    while n.lower() in used:
        suffix = f"_{i}"
        n = base[:31 - len(suffix)] + suffix
        i += 1
    used.add(n.lower())
    return n


def _cols_for(session: Session, batch_id: int, sheet_key: str, findings: list[Finding]) -> list[str]:
    sc = session.execute(
        select(SheetColumns).where(
            SheetColumns.batch_id == batch_id, SheetColumns.sheet_key == sheet_key
        )
    ).scalars().first()
    if sc and sc.columns:
        return list(sc.columns)
    # 退路：raw 鍵的有序聯集
    seen: list[str] = []
    for f in findings:
        for k in (f.raw or {}):
            if k not in seen:
                seen.append(k)
    return seen or _FALLBACK_COLS


def build_workbook(session: Session, batch_id: Optional[int] = None,
                   today: Optional[dt.date] = None,
                   department: Optional[str] = None) -> tuple[Workbook, str]:
    """回傳 (workbook, 建議檔名)。無資料則回空白帶提示。department 指定時只匯出該部門。"""
    today = today or dt.date.today()
    batch = _pick_batch(session, batch_id)
    wb = Workbook()

    # 第一張：管理摘要
    ws = wb.active
    ws.title = "管理摘要"
    _write_summary(ws, session, batch, today)
    used = {"管理摘要"}

    if batch:
        q = select(Finding).where(Finding.batch_id == batch.id)
        if department and department != "全部":
            q = q.where(Finding.department == department)
        fs = session.execute(q).scalars().all()
        by_sheet: dict[str, list[Finding]] = {}
        for f in fs:
            by_sheet.setdefault(f.sheet_key or "未分類", []).append(f)

        for sheet_key in sorted(by_sheet):
            rows = by_sheet[sheet_key]
            cols = _cols_for(session, batch.id, sheet_key, rows)
            ws = wb.create_sheet(_safe_sheet_name(sheet_key, used))
            ws.append(cols)
            for c in ws[1]:
                c.font = Font(bold=True)
            for f in rows:
                ws.append([_cell(f.raw, col) for col in cols])

    stamp = (batch.imported_at.strftime("%Y%m%d") if batch else today.strftime("%Y%m%d"))
    return wb, f"弱點原封匯出_{stamp}.xlsx"


def _write_summary(ws, session: Session, batch: Optional[ImportBatch], today: dt.date) -> None:
    ws.append(["指標", "數值"])
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="E8F5E9")
    if not batch:
        ws.append(["狀態", "尚無匯入資料"])
        return
    s = query.summary(session, today=today)
    rows = [
        ("資料批次時間", batch.imported_at.strftime("%Y-%m-%d %H:%M")),
        ("資料距今(天)", s["freshness"]["days_ago"]),
        ("未結案", s["unresolved"]),
        ("已逾期", s["overdue"]),
        ("30天內到期", s["due_soon"]),
        ("高風險(Critical+High)", s["high_risk"]),
        ("應提申請未提", s["should_apply"]),
        ("無負責人", s["gaps"]["no_owner"]),
        ("無到期日", s["gaps"]["no_due"]),
        ("已結案", s["closed"]),
        ("結案率(%)", s["close_rate"]),
    ]
    for k, v in rows:
        ws.append([k, v])


def to_bytes(wb: Workbook) -> bytes:
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ── 清單「完整匯出」：勾選／眼前那幾筆，原檔欄序、型別還原，可直接整段貼回原 Excel（2026-10-12） ──
_DATE_RE = re.compile(r"^\s*(\d{4})[/-](\d{1,2})[/-](\d{1,2})(?:[ T]00:00(?::00)?)?\s*$")
_INT_RE = re.compile(r"^\s*-?[1-9]\d{0,14}\s*$|^\s*0\s*$")
CASE_COLS = [("progress", "承辦_處理進度"), ("progress_state", "承辦_對帳狀態"),
             ("target_date", "承辦_預計完成日"), ("track_note", "承辦_追蹤備註")]


def _typed(v):
    """原值是字串時還原型別：日期字串→真日期、純數字→整數（前導 0 的保留文字，免得編號被吃掉）。"""
    if not isinstance(v, str):
        return v
    m = _DATE_RE.match(v)
    if m:
        try:
            return dt.datetime(int(m.group(1)), int(m.group(2)), int(m.group(3)))
        except ValueError:
            return v
    if _INT_RE.match(v):
        return int(v)
    return v


def build_rows_workbook(session: Session, ids: list, case_info: Optional[dict] = None) -> tuple[Workbook, str]:
    """指定的 finding（id）→ 每張來源表一頁，欄序同原檔；原檔欄之後接承辦欄。"""
    from openpyxl.styles import Alignment, PatternFill
    case_info = case_info or {}
    fs = session.execute(select(Finding).where(Finding.id.in_(list(ids) or [-1])).order_by(Finding.id)).scalars().all()
    wb = Workbook()
    wb.remove(wb.active)
    used: set = set()
    by_sheet: dict = {}
    for f in fs:
        by_sheet.setdefault(f.sheet_key or "未分類", []).append(f)
    no_wrap = Alignment(wrap_text=False, vertical="top")
    case_fill = PatternFill("solid", fgColor="E2F0D9")
    for sheet_key, rows in by_sheet.items():
        cols = _cols_for(session, rows[0].batch_id, sheet_key, rows)
        ws = wb.create_sheet(_safe_sheet_name(sheet_key, used))
        ws.append(list(cols) + [h for _, h in CASE_COLS])
        for i, c in enumerate(ws[1]):
            c.font = Font(bold=True)
            if i >= len(cols):
                c.fill = case_fill
        for f in rows:
            raw = f.raw or {}
            ci = case_info.get(f.id) or {}
            vals = [_typed(raw.get(col)) for col in cols] + [ci.get(k) for k, _ in CASE_COLS]
            ws.append(vals)
            for c in ws[ws.max_row]:
                c.alignment = no_wrap
                if isinstance(c.value, dt.datetime):
                    c.number_format = "yyyy/m/d"
        for i, col in enumerate(list(cols) + [h for _, h in CASE_COLS], start=1):
            ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = min(max(len(str(col)) * 2 + 2, 10), 24)
    if not wb.sheetnames:
        wb.create_sheet("無資料")
    return wb, "弱點清單_完整_%s.xlsx" % dt.date.today().strftime("%Y%m%d")
